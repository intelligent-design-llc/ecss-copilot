"""Excel export: Summary, Requirements, Findings, Verification Matrix, Rules."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .rules import rule_catalogue

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF4CE")
SEV_FILL = {"High": PatternFill("solid", fgColor="F8D7DA"), "Medium": PatternFill("solid", fgColor="FFE8C2"),
            "Low": PatternFill("solid", fgColor="E2EEF9")}
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
    ws.freeze_panes = "B2"
    return ws


def _finish(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def to_xlsx(result: dict) -> bytes:
    wb = Workbook()
    wb.properties.creator = "ECSS Copilot"
    wb.properties.title = "ECSS Copilot report (advisory draft)"
    s = wb.active
    s.title = "Summary"
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 90
    s.append(["ECSS Copilot - requirement quality & verification draft"])
    s["A1"].font = Font(bold=True, size=14, color="1F3A5F")
    s.append(["DISCLAIMER", result["disclaimer"]])
    s["A2"].font = Font(bold=True, color="9C0006")
    s["B2"].fill = WARN_FILL
    s["B2"].alignment = WRAP
    s.row_dimensions[2].height = 60
    sm = result["summary"]
    rows = [
        ("Tool", result["tool"]), ("Generated", result["generated_at"]), ("Source", result["source"]),
        ("Requirements extracted", sm["requirements"]), ("Total findings", sm["findings"]),
        ("Requirements with findings", sm["requirements_with_findings"]),
        ("Requirements with no findings", sm["clean_requirements"]),
        ("Findings - High", sm["by_severity"]["High"]), ("Findings - Medium", sm["by_severity"]["Medium"]),
        ("Findings - Low", sm["by_severity"]["Low"]), ("Average quality score (0-100)", sm["average_score"]),
        ("LLM pass", f"yes ({result.get('llm_model')})" if result.get("llm_used") else "no (rule-based only)"),
    ]
    for r in rows:
        s.append(list(r))
    s.append([])
    s.append(["Findings by rule", "Count"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    for k, v in sm["by_rule"].items():
        s.append([k, v])
    s.append([])
    s.append(["Draft verification methods", "Count"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    for k, v in sm["verification_methods"].items():
        s.append([k, v])

    req = _sheet(wb, "Requirements",
                 ["Req ID", "Requirement text", "Source location", "Quality score", "# High", "# Medium", "# Low",
                  "Draft method", "Suggested rewrite (LLM)", "Reviewer decision", "Reviewer comment"],
                 [16, 70, 16, 10, 8, 9, 8, 16, 60, 18, 30])
    dv = DataValidation(type="list", formula1='"Accept,Accept with changes,Reject,Clarify with customer"', allow_blank=True)
    req.add_data_validation(dv)
    for r in result["requirements"]:
        cnt = {k: sum(1 for f in r["findings"] if f["severity"] == k) for k in ("High", "Medium", "Low")}
        llm = r.get("llm") or {}
        req.append([r["id"], r["text"], r["source"], r["score"], cnt["High"], cnt["Medium"], cnt["Low"],
                    r["verification"]["primary"], llm.get("rewrite", ""), "", ""])
        dv.add(f"J{req.max_row}")
        sc = req.cell(row=req.max_row, column=4)
        sc.fill = SEV_FILL["High"] if r["score"] < 50 else SEV_FILL["Medium"] if r["score"] < 80 else PatternFill("solid", fgColor="D4EDDA")
    _finish(req)

    fnd = _sheet(wb, "Findings",
                 ["Req ID", "Rule", "Check", "Severity", "Triggered by", "Finding", "Suggested fix", "Requirement text"],
                 [16, 9, 28, 10, 22, 60, 60, 60])
    for r in result["requirements"]:
        for f in r["findings"]:
            fnd.append([f["req_id"], f["rule_id"], f["rule_name"], f["severity"], f["match"], f["message"],
                        f["suggestion"], r["text"]])
            fnd.cell(row=fnd.max_row, column=4).fill = SEV_FILL[f["severity"]]
    _finish(fnd)

    vm = _sheet(wb, "Verification Matrix",
                ["Req ID", "Requirement text", "T", "A", "R", "I", "Primary method", "Secondary method",
                 "Rationale", "Method source", "Verification level", "Stage", "Status"],
                [16, 70, 5, 5, 5, 5, 16, 16, 55, 14, 16, 14, 26])
    col = {"Test": 3, "Analysis": 4, "Review of design": 5, "Inspection": 6}
    for r in result["requirements"]:
        v = r["verification"]
        llm = r.get("llm") or {}
        primary = llm.get("verification_method") or v["primary"]
        src = "LLM" if llm.get("verification_method") else "rule-based"
        rationale = llm.get("rationale") or v["rationale"]
        row = [r["id"], r["text"], "", "", "", "", primary, v.get("secondary") or "", rationale, src,
               "TBD (e.g. unit / subsystem / system)", "TBD (e.g. QM / FM)", "DRAFT - engineer to confirm"]
        row[col[primary] - 1] = "X"
        if v.get("secondary") and v["secondary"] != primary:
            row[col[v["secondary"]] - 1] = "(x)"
        vm.append(row)
    _finish(vm)
    for row in vm.iter_rows(min_row=2, min_col=3, max_col=6):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="top")
    vm.append([])
    vm.append(["Legend", "T = Test, A = Analysis (incl. similarity), R = Review of design, I = Inspection. "
               "X = proposed primary method, (x) = proposed secondary method. All entries are drafts."])

    rc = _sheet(wb, "Rules", ["Rule ID", "Check"], [10, 50])
    for c in rule_catalogue():
        rc.append([c["rule_id"], c["name"]])
    _finish(rc)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
