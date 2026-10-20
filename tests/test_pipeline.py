import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from ecss_copilot.app import app
from ecss_copilot.core import analyse
from ecss_copilot.export import to_xlsx
from ecss_copilot.extract import extract, requirements_from_text
from ecss_copilot.models import Requirement
from ecss_copilot.verification import draft_verification

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def method(text):
    return draft_verification(Requirement("X", text)).primary


@pytest.mark.parametrize("text,expected", [
    ("The structure shall withstand the random vibration levels of [AD-4].", "Test"),
    ("The UHF transceiver shall transmit telemetry at 9.6 kbps.", "Test"),
    ("The spacecraft shall operate for a mission lifetime of 2 years.", "Analysis"),
    ("The spacecraft shall re-enter the atmosphere within 5 years after end of mission.", "Analysis"),
    ("All external surfaces shall be labelled with the serial number.", "Inspection"),
    ("The software shall comply with the project coding standard.", "Review of design"),
])
def test_verification_heuristics(text, expected):
    assert method(text) == expected


def test_text_extraction_with_ids():
    reqs = requirements_from_text("3.1 Power\nEPS-010 The EPS shall provide 5 V.\ncontinued line.\n\nEPS-020: The EPS shall provide 3.3 V.\n")
    assert [r.id for r in reqs] == ["EPS-010", "EPS-020"]
    assert reqs[0].text.endswith("continued line.")


def test_heading_with_id_like_token_is_not_requirement():
    reqs = requirements_from_text("LARK-1 System Overview\nSYS-010 The spacecraft shall have a mass below 4 kg.")
    assert [r.id for r in reqs] == ["SYS-010"]


def test_text_extraction_without_ids():
    reqs = requirements_from_text("Intro text. The EPS shall provide 5 V. The OBC shall boot in 30 s.")
    assert len(reqs) == 2 and all(r.id_generated for r in reqs)


@pytest.mark.parametrize("name", ["LARK-1_SRS_sample.docx", "LARK-1_SRS_sample.csv", "LARK-1_SRS_sample.pdf"])
def test_sample_files(name):
    reqs = extract(name, (SAMPLES / name).read_bytes())
    assert len(reqs) >= 39
    assert any(r.id == "LARK-EPS-060" for r in reqs)


def test_xlsx_export_roundtrip():
    reqs = extract("s.csv", (SAMPLES / "LARK-1_SRS_sample.csv").read_bytes())
    data = to_xlsx(analyse(reqs, "sample"))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Requirements", "Findings", "Verification Matrix", "Rules"]
    assert wb["Requirements"].max_row == len(reqs) + 1
    # the exported requirements sheet can be re-imported
    buf = io.BytesIO(data)
    again = extract("again.xlsx", buf.getvalue())
    assert len(again) >= len(reqs)


def test_api_end_to_end(monkeypatch):
    monkeypatch.delenv("LLM_BASE_URL", raising=False)
    c = TestClient(app)
    assert c.get("/api/config").json()["llm"]["enabled"] is False
    r = c.post("/api/analyze", data={"text": "SYS-010 The system shall be fast."})
    assert r.status_code == 200
    res = r.json()
    assert res["summary"]["requirements"] == 1 and res["summary"]["findings"] >= 1
    x = c.post("/api/export", json=res)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert c.post("/api/llm", json=res).status_code == 400
    assert c.post("/api/analyze", data={"text": "nothing here"}).status_code == 422


def test_wrapped_line_starting_with_id_like_token_is_continuation():
    text = ("This SRS defines the system-level requirements for the\nLARK-1 3U CubeSat, an Earth-observation technology demonstrator.\n"
            "LARK-SYS-030 The spacecraft shall operate in a 520 km sun-\nsynchronous orbit.")
    reqs = requirements_from_text(text)
    assert [r.id for r in reqs] == ["LARK-SYS-030"]
    assert "sun-synchronous" in reqs[0].text
