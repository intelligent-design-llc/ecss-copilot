"""Generate the synthetic sample spec used for the demo and tests.

SAMPLE DATA: LARK-1 is a fictional 3U CubeSat. All requirements are made up for
demonstration and contain deliberate flaws. Not derived from any real programme.
"""
import csv
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "samples"

SECTIONS = [
    ("3.1 Mission and system", [
        ("LARK-SYS-010", "The LARK-1 spacecraft shall have a total mass not exceeding 4.0 kg including margins."),
        ("LARK-SYS-020", "The spacecraft shall fit within the 3U CubeSat envelope of 100 mm x 100 mm x 340.5 mm."),
        ("LARK-SYS-030", "The spacecraft shall operate for a mission lifetime of at least 2 years in a 520 km sun-synchronous orbit."),
        ("LARK-SYS-040", "The spacecraft shall be easy to integrate and shall be compatible with all commercial deployers."),
        ("LARK-SYS-050", "The spacecraft shall re-enter the atmosphere within 5 years after end of mission."),
        ("LARK-SYS-060", "The system should be designed with high reliability."),
        ("LARK-SYS-070", "The spacecraft shall comply with the launch provider interface control document [AD-3]."),
    ]),
    ("3.2 Electrical power subsystem (EPS)", [
        ("LARK-EPS-010", "The EPS shall provide a regulated 5 V ± 0.1 V bus with a maximum ripple of 50 mV peak-to-peak."),
        ("LARK-EPS-020", "The EPS shall provide adequate power to all subsystems in all modes."),
        ("LARK-EPS-030", "The battery shall have a capacity of 40 and shall be charged by the solar arrays."),
        ("LARK-EPS-040", "The EPS shall generate at least 7.5 W orbit-average power at end of life."),
        ("LARK-EPS-050", "The EPS shall switch the spacecraft to safe mode within 2 s when the battery voltage drops below 6.4 V."),
        ("LARK-EPS-060", "Shall protect the battery against overcharge and/or deep discharge."),
        ("LARK-EPS-070", "The EPS shall provide 3 switchable power lines to the payload, current limited to TBD mA."),
    ]),
    ("3.3 On-board computer and software (OBC)", [
        ("LARK-OBC-010", "The OBC shall boot into safe mode within 30 s after power-on."),
        ("LARK-OBC-020", "The on-board software shall be user-friendly and fast."),
        ("LARK-OBC-030", "The OBC shall store at least 2 GB of payload data."),
        ("LARK-OBC-040", "The on-board software shall never crash under any conditions."),
        ("LARK-OBC-050", "Telecommands shall be time-tagged and executed as appropriate."),
        ("LARK-OBC-060", "The OBC shall reject telecommands with an invalid CRC and shall increment the rejected-command counter by 1."),
        ("LARK-OBC-070", "The on-board software shall be written in C and comply with the project coding standard [AD-5]."),
        ("LARK-OBC-080", "The watchdog shall reset the OBC if it is not serviced within 10."),
    ]),
    ("3.4 Communications (COM)", [
        ("LARK-COM-010", "The UHF transceiver shall transmit telemetry at a data rate of 9.6 kbps."),
        ("LARK-COM-020", "The S-band downlink shall provide a link margin of at least 3 dB at 10 deg elevation."),
        ("LARK-COM-030", "The communication system will support various ground stations, etc."),
        ("LARK-COM-040", "The UHF antenna shall deploy no earlier than 30 min after separation from the deployer."),
        ("LARK-COM-050", "The S-band transmitter shall not transmit before the ground has sent the TX-enable command."),
        ("LARK-COM-060", "The spacecraft shall maintain contact with the ground as far as possible."),
    ]),
    ("3.5 Attitude determination and control (ADCS)", [
        ("LARK-ADC-010", "The ADCS shall point the payload boresight with an accuracy better than 1 deg (3 sigma)."),
        ("LARK-ADC-020", "The ADCS shall detumble the spacecraft from 10 deg/s to below 0.5 deg/s within 3 orbits."),
        ("LARK-ADC-030", "The ADCS shall provide sufficient pointing stability for imaging, TBC."),
        ("LARK-ADC-040", "The ADCS shall use a Kalman filter for attitude estimation."),
    ]),
    ("3.6 Payload, structure and environment", [
        ("LARK-PLD-010", "The payload camera shall acquire images with a ground sampling distance of 30 m from 520 km altitude."),
        ("LARK-STR-010", "The structure shall withstand the qualification random vibration levels of [AD-4] without damage."),
        ("LARK-STR-020", "The first natural frequency of the spacecraft shall be above 100 Hz."),
        ("LARK-STR-030", "All external surfaces shall be labelled with the spacecraft serial number."),
        ("LARK-THR-010", "The thermal control shall keep the battery between 0 °C and 40 °C in all operational modes."),
        ("LARK-THR-020", "The thermal design shall keep it within limits."),
        ("LARK-STR-040", "The spacecraft shall be provided with a remove-before-flight pin accessible after integration into the deployer, it shall be colored red, and it shall be included in the delivered hardware with documentation describing how it is installed and removed by the integration team at the launch site."),
    ]),
]

INTRO = [
    "SAMPLE DATA - FICTIONAL MISSION. This document is synthetic and was created to demonstrate ECSS Copilot. "
    "LARK-1 is not a real spacecraft, and the requirements contain deliberate flaws.",
    "1 Scope: this System Requirements Specification (SRS) defines the system-level requirements for the LARK-1 "
    "3U CubeSat, an Earth-observation technology demonstrator.",
    "2 Applicable documents: [AD-3] Launch provider ICD, [AD-4] Launch environment specification, [AD-5] Software coding standard.",
]


def main():
    OUT.mkdir(exist_ok=True)
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("LARK-1 System Requirements Specification (SAMPLE)", 0)
    p = d.add_paragraph()
    r = p.add_run("Doc. LARK-SRS-001, Issue 0.3 (draft) - SAMPLE DATA, fictional mission, synthetic content")
    r.bold = True
    r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    for para in INTRO:
        d.add_paragraph(para)
    d.add_heading("3 Requirements", 1)
    rows_csv = []
    for title, reqs in SECTIONS:
        d.add_heading(title, 2)
        t = d.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Req ID", "Requirement text", "Rationale"
        for rid, text in reqs:
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text = rid, text, ""
            rows_csv.append((rid, text, title))
    d.add_heading("4 Requirements without ID (draft notes from the review)", 1)
    d.add_paragraph("The ground segment shall archive all telemetry. The mission control software should display alarms. "
                    "The ground segment shall archive all telemetry.")
    cp = d.core_properties
    cp.author = cp.last_modified_by = "Intelligent Design LLC"
    cp.title = "LARK-1 System Requirements Specification (SAMPLE DATA)"
    cp.comments = "Synthetic sample data for ECSS Copilot. Fictional mission."
    cp.created = cp.modified = datetime(2026, 10, 1, tzinfo=timezone.utc)
    d.save(OUT / "LARK-1_SRS_sample.docx")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:  # optional PDF variant
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", str(OUT / "LARK-1_SRS_sample.docx"),
                        "--outdir", str(OUT)], check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    with open(OUT / "LARK-1_SRS_sample.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["ID", "Requirement", "Section"])
        w.writerows(rows_csv)
    print("wrote", OUT / "LARK-1_SRS_sample.docx", "and .csv;", len(rows_csv), "table requirements")


if __name__ == "__main__":
    main()
